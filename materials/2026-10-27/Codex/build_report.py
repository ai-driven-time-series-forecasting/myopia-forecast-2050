"""Build the requested short Word note and a detailed reproducibility guide."""
import argparse
import csv
import hashlib
import json
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--out',required=True);ap.add_argument('--input',required=True)
    args=ap.parse_args();out=Path(args.out)
    a=json.loads((out/'run_manifest.json').read_text(encoding='utf-8'))
    metrics=list(csv.DictReader((out/'metrics.csv').open(encoding='utf-8-sig')))
    mean=a['cv_mean'];h=a['balanced_historical_comparison']
    baseline=list(csv.DictReader((out/'baseline_metrics.csv').open(encoding='utf-8-sig')))
    naive=next(x for x in baseline if x['fold']=='mean' and x['model']=='last_observation')
    ablation=next(x for x in baseline if x['fold']=='mean' and x['model']=='same_model_without_spatial')
    prov=list(csv.DictReader((out/'province_standardized_trends.csv').open(encoding='utf-8-sig')))
    future2050=[x for x in prov if x['year']=='2050']
    scenarios={k:sum(float(x[k]) for x in future2050)/31 for k in ['phi090','predicted','phi100']}
    note=[
        '本分析由Codex（基于GPT-6，精确版本未披露）组织，数值由Python计算。采用computed_prevalence，剔除263行无效记录；同panel同年多项研究按病例数与样本量合并，得到1791个panel年份单元。',
        '模型在Logit尺度控制年级、城乡、验光方式及交互项，估计省份截距和趋势；用省际距离构建四近邻图，以空间惩罚共享信息。每年趋势按0.95衰减，逆变换保证比例介于0与1。',
        f'十折按年份向前验证，训练年严格早于测试年，参数仅在训练集内部选择。十折均值：MSE {mean["MSE"]:.4f}、RMSE {mean["RMSE"]:.4f}、MAE {mean["MAE"]:.4f}、MAPE {mean["MAPE"]:.2f}%、R² {mean["R2"]:.3f}。误差仍大，仅适合探索趋势，精度尚不令人满意。',
        f'1995与2019年均有观测的224个panel，其等权均值由{h["mean_1995"]:.1%}升至{h["mean_2019"]:.1%}。模型按四年级组×城乡等权、验光方式3标准化，31省均值由2022年{a["standardized_mean_2022"]:.1%}升至2050年{a["standardized_mean_2050"]:.1%}，增速趋缓；这不是人口加权全国率。',
        '2022年仅3省有观测，59个panel无有效历史，预测依赖信息共享。17年远期回测RMSE约31.0个百分点；短期精度不能验证2050年精度。政策、教育和户外活动变化未建模。图中带状范围为趋势情景敏感性，非置信区间；稀疏地区及远期结果应谨慎使用。'
    ]
    doc=Document();sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2)
    normal=doc.styles['Normal'];normal.font.name='Microsoft YaHei';normal.font.size=Pt(10.5)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    normal.paragraph_format.line_spacing=1.35
    normal.paragraph_format.space_after=Pt(8)
    normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    for style in doc.styles:
        for border in style._element.xpath('.//*[local-name()="pBdr"]'):
            border.getparent().remove(border)
    for name in ['Title','Heading 1','Heading 2']:
        st=doc.styles[name];st.font.name='Microsoft YaHei';st._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
        st.font.color.rgb=RGBColor.from_string('214D74')
    doc.styles['Title'].font.size=Pt(22)
    doc.styles['Heading 1'].font.size=Pt(12)
    p=doc.add_paragraph('青少年近视患病率预测说明',style='Title')
    p.paragraph_format.space_after=Pt(4)
    p=doc.add_paragraph('Codex · 省级面板 · 2023–2050')
    p.runs[0].font.size=Pt(10);p.runs[0].font.color.rgb=RGBColor.from_string('647482')
    doc.add_heading('方法与验证',level=1)
    for t in note[:3]:doc.add_paragraph(t)
    doc.add_heading('趋势解读与不确定性',level=1)
    for t in note[3:]:doc.add_paragraph(t)
    p=doc.add_paragraph('图示：灰点及连线为调查样本合并比例，蓝实线为历史模型拟合，蓝虚线为未来预测。两者构成口径不同；没有将缺失年份补为实测值。')
    p.paragraph_format.space_after=Pt(6)
    p=doc.add_paragraph('方法参考：scikit-learn 时间序列交叉验证文档；Hyndman & Athanasopoulos, Forecasting: Principles and Practice (3rd ed.), §8.2（趋势衰减思想）。')
    p.paragraph_format.space_after=Pt(0)
    for r in p.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('647482')
    doc.core_properties.title='青少年近视患病率预测说明'
    doc.core_properties.author='Codex'
    doc.core_properties.subject='空间正则化面板预测、十折时间交叉验证与2050年情景预测'
    doc.save(out/'预测说明_Codex.docx')
    print('Main note character count:',sum(len(t) for t in note),flush=True)
    fold_table='\n'.join('| '+ ' | '.join(str(r[k]) for k in ['fold','train_end','test_start','test_end','n_train','n_test','selected_lambda'])+' |' for r in metrics[:-1])
    md=f'''# 省级青少年近视率预测：方法、口径与复现

## 交付与总体结果

- `metrics.csv`：十折逐折指标，最后一行 `fold=mean` 为十个折指标的算术均值。
- `predictions_future.csv`：534 × 28 = 14,952行；仅含要求的六列。unique_id保留原始method后缀。
- `province_forecasts_300dpi.png`：一张31省小多图，5700×5700像素，300 dpi。
- `预测说明_Codex.docx`：简要研究说明；模型名称为Codex（基于GPT-6），未冒充环境未提供的精确版本。

十折平均：MSE={mean['MSE']:.8f}，RMSE={mean['RMSE']:.8f}，MAE={mean['MAE']:.8f}，MAPE={mean['MAPE']:.4f}%，R²={mean['R2']:.6f}。
RMSE与MAE在0–1比例尺度，乘100为百分点。MSE的单位是比例平方，R²无单位；MAPE已乘100。

## 数据与清洗

来源为用户提供的06_myopia.zip，观测文件SHA256：`{a['input_sha256']}`。
2096行观测，31省，1995–2022中仅25个实际观测年份。263行目标缺失/计数无效，不作零值填补；其中259行标记零分母。
1833行有效观测经同unique_id、同year的病例数/样本量合并，形成1791个panel年份单元，原始475个panel有有效历史，59个panel无有效历史。
42行是同panel同年的重复研究单元，未发现病例数、样本量和其余键完全相同的内容重复。合并不证明研究样本独立，仍可能存在未知重叠。
目标严格为computed_prevalence；已有y通常是reported_prevalence，1767行与目标相差超过1e-6；不将y直接用于拟合。
22行rate_disagreement为真，保留经计数验证的computed_prevalence，不改用报告率。
所有原始行与输入压缩包不作修改；year为四位整数，ds原格式YYYY-MM-DD，年份一致性已检查。

`cleaned_panel_year.csv`保存同panel同年的计数合并结果。每个合并后的单元在回归损失和五项指标中等权。
大调查不能被机械视为无异质性的数十万个独立样本，因此回归不直接以全部sample_size作精度权重。
sample_size仅用于同panel同年合并和原始调查比例的描述性汇总；病例数、样本量、报告率、split不作为预测特征。

## 模型与空间特征

这是本次实现的空间正则化Logit面板回归，不是由大模型逐项猜数，也不是标准Holt模型或因果模型。

z=logit(clip(p,1e-5,1-1e-5))；最小化 Σ(z-Xβ)² + βᵀRβ。
回归有201项系数：截距、每十年趋势、年级哑变量、城乡、验光方法、年级×城乡、年级×时间、城乡×时间，以及省份截距、省份趋势、省份×年级、省份×城乡偏移。
参考年级为1–3年级，参考居住地为农村，参考方法为3。类别字典是固定的已知学段、方法与行政区，不是从测试目标学习。
训练时间中心为该折最后一个实际训练年，避免用测试年设定中心。训练只使用实际年份，不人为构造逐年目标。

使用提供的省际距离矩阵：每省取4个近邻，w=exp(-d/800)，对称化并使平均节点度为1；L=diag(W1)-W。
源文件未注明距离单位，因此仅将其称为提供的距离数值尺度，不宣称精确交通距离。
R中省截距加3λL、省趋势加5λL。省截距的独立收缩强度λ，省趋势10λ，省份×年级及省份×城乡3λ；全局趋势和相关交互、方法系数弱正则化0.1。
其余全局系数的1e-5惩罚保证可识别性。全部惩罚与模型系数可在脚本及model_coefficients.csv检查。
λ候选为0.3、3、30，各折只在训练集内最多3个向前两观测年块上，以平均MSE选择。训练年不足时固定λ=3。
最终用截至2022年的所有有效数据重新训练，最终λ={a['final_lambda']}。
sunshine-hours-rank.csv没有观测日期、时变暴露或未来取值说明，本次不使用，以免混淆静态排名与可外推的因果因素。

未来h年使用有效时间Σ(k=1..h)φ^k，主情景φ=0.95；历史部分保留真实线性时间。
这是主动声明的远期趋势衰减假设，未将其冒充由短期回测充分估计的参数。
逆Logit将预测映射回0–1。无有效历史的59个panel借助省份、邻省和分层变量估计，不能理解为这些panel获得了独立的时间序列验证。

## 十折验证与无泄漏控制

按25个实际年份排序，最早5个年份构成初始训练集，随后20个年份分成10个前向测试块，每块两个实际观测年。
例如2009年没有观测，所以第4折测试2008与2010，不宣称所有折具有相同日历跨度。
同一年所有省份、年级、城乡和方法全部留在同一侧，绝不行级随机打乱。
允许较早折的测试数据在较晚折中成为历史训练数据，这正是扩展窗口；绝不使用当前折或更晚的测试数据训练当前折。
所有目标处理仅在同年同panel内完成，不跨年插值；每折模型和参数选择重新拟合。没有将外层十折的结果用于选择最终λ。
原文件的train/validation/test标签不沿用，因为本任务明确要求重新作十折时间交叉验证。

|折|训练截至|测试起年|测试止年|训练单元|测试单元|λ|
|---|---|---|---|---|---|---|
{fold_table}

MSE=mean((y-p)²)，RMSE=sqrt(MSE)，MAE=mean(abs(y-p))，MAPE=100×mean(abs((y-p)/y))，R²=1-Σ(y-p)²/Σ(y-mean(y))²。
本数据没有零值有效目标；通用代码遇零值时仅将其排除MAPE分母并记录数量，不把零值改成任意小数。
mean行逐指标等权平均10个折。不能把sqrt(平均MSE)当作平均RMSE，也不能用合并全部测试点的R²替代平均R²。
run_manifest.json另列cv_pooled，仅作诊断，不能与任务要求的cv_mean混用。测试量很小的年份仍在折均值中占同等权重，MAPE又受低患病率放大。
cv_predictions.csv可逐条核对每折真实目标、预测、截止年与panel。

## 基准、空间消融与远期检验

沿用最近一次panel观测（新panel回退为训练期年级×城乡均值）的十折平均RMSE={float(naive['RMSE']):.6f}。
本模型相对这一基准的RMSE改善约{100*(1-mean['RMSE']/float(naive['RMSE'])):.1f}%。
去掉图拉普拉斯、保留同一折λ与其他特征的消融平均RMSE={float(ablation['RMSE']):.6f}，空间项改善很小；不据此宣称显著空间预测优势，消融也不是独立调参的最优非空间模型。

long_horizon_diagnostics.csv补充冻结训练截至2005年的9/14/17年回测；参数只由截至2005年内部验证选择，未用于修改主模型。
2022年17年预测RMSE约0.3102、R²约-0.088，且只有11个测试单元。这提示远期可靠性有限，不能用一至三年的外层验证替代28年验证。

## 省级图、标准化与解读

省级蓝线为固定四年级组×城乡8层等权、统一method_code=3的模型均值。灰点及灰连线为当年原始调查病例总和/样本量总和，其样本构成随年改变，不能直接把灰点对蓝线的偏离当作模型误差。
1995–2022蓝实线是最终模型回顾拟合，2023–2050蓝虚线是外推；无观测年份未伪造实测点。西藏仅2017–2020有有效观测，此前蓝线属于跨省共享信息的历史推算。
图包含实际观测与连续拟合/预测，但这些历史拟合绝不用于报告训练内精度。
预测表仍按原始534个方法特定panel逐项输出；图中的标准化省级率不能用预测表按省直接无权平均来复现。图的完整数值保存于province_standardized_trends.csv。

可比224个panel的观测等权均值，1995年{h['mean_1995']:.1%}，2019年{h['mean_2019']:.1%}，上升{h['change_pp']:.1f}个百分点；它描述所覆盖调查panel，不是全国人口加权流行率。
31省等权、固定8层的模型均值，1995年{a['standardized_mean_1995']:.1%}、2022年{a['standardized_mean_2022']:.1%}、2050年{a['standardized_mean_2050']:.1%}。
2050年各省主情景标准化率约{a['minimum_2050']:.1%}–{a['maximum_2050']:.1%}，省份数值见province_2050_summary.csv。不能将此范围解释为预测置信区间。
2050年31省均值在φ=0.90、0.95、1.00三种假设下分别为{scenarios['phi090']:.1%}、{scenarios['predicted']:.1%}、{scenarios['phi100']:.1%}。
金色范围是三个趋势假设逐省逐年的最小值到最大值，不覆盖抽样误差、参数误差、干预冲击，也没有80%/95%概率含义。
没有可靠的远期校准样本，故不制造看似精确的概率预测区间。政策、户外活动、教育结构、诊断定义、人口组成和调查异质性都可能改变未来走势。

## 复现

需要Python、numpy、scipy、pandas、matplotlib及python-docx；requirements.txt给出最低兼容要求。
在解压完整交付包的目录执行：

```powershell
python model_forecast.py --input input/06_myopia.zip --out reproduced
python build_report.py --input input/06_myopia.zip --out reproduced
```

模型计算本次使用Python {a['python']}，numpy {a['numpy']}，pandas {a['pandas']}，matplotlib {a['matplotlib']}。
图中文字字体优先使用Microsoft YaHei，其他平台需安装可用中文字体。
run_manifest.json保存输入哈希、质量统计、参数与检查；nested_cv_selection.json保存外层每折的内部选择记录。

## 参考资料

1. scikit-learn, TimeSeriesSplit： https://scikit-learn.org/stable/modules/generated/sklearn.model_selection.TimeSeriesSplit.html
2. Hyndman & Athanasopoulos, Forecasting: Principles and Practice (3rd ed.), §8.2： https://otexts.com/fpp3/holt.html
第2项支持衰减趋势这一预测思想；本次实现是Logit空间回归，不声称执行了该书的完整Holt估计程序。
'''
    (out/'methodology.md').write_text(md,encoding='utf-8')
    (out/'requirements.txt').write_text('numpy>=1.26\nscipy>=1.11\npandas>=2.1\nmatplotlib>=3.8\npython-docx>=1.1\n',encoding='utf-8')
    here=Path(__file__).resolve().parent
    for fn in ['model_forecast.py','build_report.py']:
        if (here/fn).resolve() != (out/fn).resolve():shutil.copy2(here/fn,out/fn)
    # Package is created by final verification, after the Word rendering check.
    print('Word note and methodology written.',flush=True)


if __name__=='__main__':main()
