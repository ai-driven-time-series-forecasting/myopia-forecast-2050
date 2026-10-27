# -*- coding: utf-8 -*-
"""生成《预测说明_豆包.docx》。"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT = r"C:\Users\echo_\OneDrive\Desktop\近视率预测_2050\预测说明_豆包.docx"

doc = Document()

normal = doc.styles["Normal"]
normal.font.name = "SimSun"
normal.font.size = Pt(11)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def set_run(run, font_cn="宋体", size=11, bold=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_cn)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def heading(text, size=12.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run(r, font_cn="黑体", size=size, bold=True)
    return p


def body(text, indent=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.4
    pf.space_after = Pt(4)
    if indent:
        pf.first_line_indent = Pt(22)
    r = p.add_run(text)
    set_run(r, size=11)
    return p

t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("青少年近视率面板预测说明")
set_run(r, font_cn="黑体", size=16, bold=True)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("模型：豆包（Doubao）大模型    数据：1995–2022 年各省面板    预测期：2023–2050 年")
set_run(r, size=10.5, color=(90, 90, 90))

heading("一、方法与建模思路")
body("本预测由豆包（Doubao）大模型完成。数据为 1995–2022 年 31 个省份、4 个年级段、城乡、3 种验光"
     "方法的近视患病率面板，剔除零样本无效记录后共 1833 条有效观测、475 个面板，目标变量为 "
     "computed_prevalence。")
body("先对患病率做 logit 变换再建模，保证预测值在 0–1 之间且自然饱和。特征兼顾时间与空间：时间上含"
     "年份趋势及面板、省份、年级×城乡、全国等分层级的历史均值与线性斜率（仅用当年之前数据）；空间上"
     "含由省际距离矩阵提取的二维坐标、反距离加权空间滞后、最近五省均值和日照时数排名；另设省份、年级、"
     "城乡、验光方法哑变量及省份×时间趋势。模型为 L2 正则线性回归（Ridge，系数经网格验证），同条件下"
     "优于面板线性趋势、GBDT 及 Ridge+GBDT 混合模型。验证采用十折滚动起点交叉验证：2013–2022 年依次"
     "作为测试年，训练仅用更早年份，不随机打乱、无未来信息泄露。")

heading("二、对 2050 年预测结果的解读")
body("历史上全国（省际等权）平均患病率约由 1995 年的 39% 升至 2014 年的 56%，2019 年约 53%，二十多年"
     "上升约 14 个百分点，低年级段上升最快、城乡差异明显。预测将继续上升：2023 年约 59%、2030 年约"
     " 64%、2040 年约 71%、2050 年约 78%，但随饱和逐步放缓；2050 年各省约 62%（新疆）至 90%（天津），"
     "高中年级段约 93%，空间与人群差异仍将存在。")

heading("三、模型精度与远期预测的不确定性")
body("十折交叉验证指标（按折计算后取均值）如下表；整体精度可用，大样本波次年份表现较好，2015、2022 等"
     "仅含少量观测的年份误差明显偏大。")

table = doc.add_table(rows=6, cols=3)
table.style = "Table Grid"
data = [
    ("指标", "十折宏平均", "池化（全部测试样本）"),
    ("MSE", "0.024", "0.017"),
    ("RMSE", "0.150", "0.130"),
    ("MAE", "0.116", "0.099"),
    ("MAPE", "41.5%", "27.3%"),
    ("R²", "0.48", "0.70"),
]
for i, row in enumerate(data):
    for j, val in enumerate(row):
        cell = table.cell(i, j)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = p.add_run(val)
        set_run(rr, size=10.5, bold=(i == 0), font_cn="黑体" if i == 0 else "宋体")

body("但交叉验证只能检验约一年步长，而 2050 年是 28 年的远期外推：模型的线性 logit 趋势隐含患病率"
     "持续上升，无法内生反映 2018 年以来综合防控政策、用眼行为变化与人口流动，稀疏面板还依赖层级与"
     "邻省借值。因此远期结果应视为趋势性情景而非精确点预测，近五年参考价值较高，2040 年后不确定性"
     "显著增大，建议随新调查数据滚动更新。")

doc.save(OUT)
txt = "".join(p.text for p in doc.paragraphs)
print("saved:", OUT)
print("paragraph chars:", len(txt))
